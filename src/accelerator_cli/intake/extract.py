"""Adapters around the existing document extractor plus simple CSV intake."""
from __future__ import annotations

import csv
import importlib.util
import os
import pathlib
import sys
from typing import Any

from ..repository import RepositoryContext

DEFAULT_MAX_SOURCE_BYTES = 50 * 1024 * 1024


def extract_document(
    context: RepositoryContext,
    path: pathlib.Path,
) -> dict[str, Any]:
    resolved = path.expanduser().resolve()
    if not resolved.exists() or not resolved.is_file():
        raise FileNotFoundError(f"Document does not exist: {resolved}")
    max_bytes = int(
        os.getenv("ACCELERATOR_MAX_SOURCE_BYTES", str(DEFAULT_MAX_SOURCE_BYTES))
    )
    size = resolved.stat().st_size
    if size > max_bytes:
        raise ValueError(
            f"Document is {size} bytes, above the configured {max_bytes}-byte "
            "intake limit."
        )
    suffix = resolved.suffix.lower()
    if suffix == ".csv":
        return _extract_csv(resolved)
    if suffix in {".xlsx", ".xlsm"}:
        return _extract_xlsx(resolved)
    if suffix == ".pptx":
        return _extract_pptx(resolved)
    extractor = _load_extractor(context)
    functions = {
        ".md": extractor.extract_markdown,
        ".txt": extractor.extract_text,
        ".docx": extractor.extract_docx,
        ".pdf": extractor.extract_pdf,
    }
    function = functions.get(suffix)
    if function is None:
        raise ValueError(
            f"Unsupported document type {suffix!r}. Supported: "
            ".md, .txt, .csv, .docx, .pdf, .pptx, .xlsx, .xlsm."
        )
    try:
        value = function(resolved)
    except SystemExit as exc:
        raise RuntimeError(
            f"Document extraction failed for {resolved.name} (exit {exc.code})."
        ) from exc
    if not isinstance(value, dict):
        raise RuntimeError(f"Extractor returned an invalid result for {resolved.name}.")
    return value


def _load_extractor(context: RepositoryContext) -> Any:
    path = context.root / "scripts" / "extract-brief-from-doc.py"
    spec = importlib.util.spec_from_file_location("accelerator_document_extractor", path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Could not load document extractor: {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _extract_csv(path: pathlib.Path) -> dict[str, Any]:
    chunks: list[dict[str, Any]] = []
    total_chars = 0
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        reader = csv.reader(handle)
        for index, row in enumerate(reader):
            text = " | ".join(cell.strip() for cell in row).strip()
            if not text:
                continue
            total_chars += len(text)
            chunks.append(
                {
                    "chunk_id": f"c{index:03d}",
                    "type": "heading" if index == 0 else "paragraph",
                    "level": 1 if index == 0 else None,
                    "text": f"[CSV row] {text}",
                    "page": None,
                }
            )
    if not chunks:
        raise ValueError(f"CSV contains no extractable rows: {path}")
    return {
        "format": "csv",
        "source_path": str(path),
        "total_chars": total_chars,
        "headings_index": [
            {
                "level": 1,
                "text": chunks[0]["text"],
                "chunk_id": chunks[0]["chunk_id"],
                "page": None,
            }
        ],
        "chunks": chunks,
    }


def _extract_xlsx(path: pathlib.Path) -> dict[str, Any]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError(
            'Install spreadsheet intake with `python -m pip install -e "."`.'
        ) from exc
    workbook = openpyxl.load_workbook(
        path,
        read_only=True,
        data_only=True,
    )
    chunks: list[dict[str, Any]] = []
    headings: list[dict[str, Any]] = []
    total_chars = 0
    index = 0
    try:
        for sheet in workbook.worksheets:
            heading_id = f"c{index:03d}"
            heading_text = f"Worksheet: {sheet.title}"
            chunks.append(
                {
                    "chunk_id": heading_id,
                    "type": "heading",
                    "level": 1,
                    "text": heading_text,
                    "page": None,
                }
            )
            headings.append(
                {
                    "level": 1,
                    "text": heading_text,
                    "chunk_id": heading_id,
                    "page": None,
                }
            )
            index += 1
            for row in sheet.iter_rows(values_only=True):
                values = [str(value).strip() for value in row if value not in (None, "")]
                if not values:
                    continue
                text = "[worksheet row] " + " | ".join(values)
                total_chars += len(text)
                chunks.append(
                    {
                        "chunk_id": f"c{index:03d}",
                        "type": "paragraph",
                        "level": None,
                        "text": text,
                        "page": None,
                    }
                )
                index += 1
    finally:
        workbook.close()
    if total_chars == 0:
        raise ValueError(f"Workbook contains no extractable cell values: {path}")
    return {
        "format": "xlsx",
        "source_path": str(path),
        "total_chars": total_chars,
        "headings_index": headings,
        "chunks": chunks,
    }


def _extract_pptx(path: pathlib.Path) -> dict[str, Any]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError(
            'Install PowerPoint intake with `python -m pip install -e "."`.'
        ) from exc
    presentation = Presentation(str(path))
    chunks: list[dict[str, Any]] = []
    headings: list[dict[str, Any]] = []
    total_chars = 0
    index = 0
    for slide_number, slide in enumerate(presentation.slides, start=1):
        title = ""
        if slide.shapes.title is not None:
            title = (slide.shapes.title.text or "").strip()
        heading_text = title or f"Slide {slide_number}"
        heading_id = f"c{index:03d}"
        chunks.append(
            {
                "chunk_id": heading_id,
                "type": "heading",
                "level": 1,
                "text": heading_text,
                "page": slide_number,
            }
        )
        headings.append(
            {
                "level": 1,
                "text": heading_text,
                "chunk_id": heading_id,
                "page": slide_number,
            }
        )
        index += 1
        seen = {heading_text}
        for shape in slide.shapes:
            texts: list[str] = []
            if getattr(shape, "has_text_frame", False):
                text = str(getattr(shape, "text", "") or "").strip()
                if text:
                    texts.append(text)
            if getattr(shape, "has_table", False):
                table = getattr(shape, "table", None)
                if table is None:
                    continue
                for row in table.rows:
                    row_text = " | ".join(
                        (cell.text or "").strip() for cell in row.cells
                    ).strip(" |")
                    if row_text:
                        texts.append(f"[table row] {row_text}")
            for text in texts:
                if text in seen:
                    continue
                seen.add(text)
                total_chars += len(text)
                chunks.append(
                    {
                        "chunk_id": f"c{index:03d}",
                        "type": "paragraph",
                        "level": None,
                        "text": text,
                        "page": slide_number,
                    }
                )
                index += 1
    if total_chars == 0:
        raise ValueError(f"Presentation contains no extractable text: {path}")
    return {
        "format": "pptx",
        "source_path": str(path),
        "total_chars": total_chars,
        "headings_index": headings,
        "chunks": chunks,
    }
