from __future__ import annotations

import pathlib

import pytest
from openpyxl import Workbook
from pptx import Presentation

from src.accelerator_cli.intake.extract import extract_document
from src.accelerator_cli.repository import RepositoryContext


def _context(tmp_path: pathlib.Path) -> RepositoryContext:
    (tmp_path / ".git").mkdir()
    return RepositoryContext(tmp_path)


def test_xlsx_intake_extracts_worksheet_rows(tmp_path: pathlib.Path) -> None:
    path = tmp_path / "requirements.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Requirements"
    sheet.append(["ID", "Requirement"])
    sheet.append(["R1", "Use managed identity"])
    workbook.save(path)

    result = extract_document(_context(tmp_path), path)

    assert result["format"] == "xlsx"
    assert any("Use managed identity" in chunk["text"] for chunk in result["chunks"])


def test_pptx_intake_extracts_slide_text(tmp_path: pathlib.Path) -> None:
    path = tmp_path / "requirements.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Security requirements"
    slide.placeholders[1].text = "Use private endpoints for regulated data."
    presentation.save(path)

    result = extract_document(_context(tmp_path), path)

    assert result["format"] == "pptx"
    assert any("private endpoints" in chunk["text"] for chunk in result["chunks"])


def test_intake_rejects_oversized_documents_before_parsing(
    tmp_path: pathlib.Path,
    monkeypatch,
) -> None:
    path = tmp_path / "oversized.txt"
    path.write_bytes(b"x" * 11)
    monkeypatch.setenv("ACCELERATOR_MAX_SOURCE_BYTES", "10")

    with pytest.raises(ValueError, match="above the configured"):
        extract_document(_context(tmp_path), path)
